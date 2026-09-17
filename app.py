import io
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import streamlit as st

st.set_page_config(page_title="エクセルリスト生成ツール", page_icon="📄", layout="wide")

# 指定のカラーコード（背景: #355E3B, ボックス: #FBFDE4, 文字: #895129）を適用したCSS
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Zen+Kaku+Gothic+New:wght@700;900&display=swap');

    /* マットな単色背景: #355E3B */
    .stApp {
        background-color: #355E3B !important;
        background-image: none !important;
    }
    
    /* メインカード: #FBFDE4 */
    .main-card {
        background-color: #FBFDE4 !important;
        border: 2px solid #895129 !important;
        border-radius: 24px;
        padding: 48px 40px;
        box-shadow: 0 15px 35px rgba(0, 0, 0, 0.25);
        color: #895129 !important;
        text-align: center;
        max-width: 640px;
        margin: 20px auto 30px auto;
        position: relative;
    }

    /* バッジ: #895129 背景 / #FBFDE4 文字 */
    .badge-theme {
        background-color: #895129 !important;
        color: #FBFDE4 !important;
        padding: 7px 22px;
        border-radius: 30px;
        font-size: 11px;
        font-weight: 800;
        letter-spacing: 2.5px;
        text-transform: uppercase;
        display: inline-block;
        margin-bottom: 18px;
    }

    /* タイトル: #895129 */
    h1.main-title {
        font-family: 'Zen Kaku Gothic New', 'Hiragino Sans', sans-serif;
        color: #895129 !important;
        font-size: 32px;
        font-weight: 900;
        margin: 0 0 16px 0;
        letter-spacing: -0.5px;
        -webkit-text-fill-color: #895129 !important;
    }

    /* 説明文: #895129 */
    p.sub-desc {
        color: #895129 !important;
        font-size: 13.5px;
        font-weight: 700;
        line-height: 1.65;
        margin-bottom: 0;
    }

    /* ボタン: #895129 背景 */
    .stButton > button {
        background-color: #895129 !important;
        color: #FBFDE4 !important;
        border: none !important;
        padding: 16px 36px !important;
        border-radius: 50px !important;
        font-family: 'Zen Kaku Gothic New', sans-serif !important;
        font-size: 17px !important;
        font-weight: 900 !important;
        letter-spacing: 1.5px !important;
        box-shadow: 0 8px 20px rgba(0, 0, 0, 0.2) !important;
        transition: all 0.3s ease !important;
        width: 100% !important;
    }

    .stButton > button:hover {
        transform: translateY(-2px) !important;
        background-color: #704020 !important;
    }

    /* ファイルアップローダー: #FBFDE4 ボックス化 */
    div[data-testid="stFileUploader"] {
        background-color: #FBFDE4 !important;
        border: 2px dashed #895129 !important;
        border-radius: 20px;
        padding: 12px;
    }

    /* アップローダー内部の文字色統一 */
    div[data-testid="stFileUploader"] * {
        color: #895129 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="main-card">
    <span class="badge-theme">PROFESSIONAL TOOL</span>
    <h1 class="main-title">エクセルリスト生成</h1>
    <p class="sub-desc">ファイルをアップロードするだけで、デザイン整形・文字装飾ルール適用・Scope別タブ分割を全自動で行います。</p>
</div>
""", unsafe_allow_html=True)

uploaded_file = st.file_uploader("", type=["xlsx", "xls", "csv"])

def clean_str(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ["nan", "none", "null"]:
        return ""
    return s

def sanitize_sheet_name(name):
    clean = re.sub(r'[:\\/*?\[\]]', '-', str(name)).strip()
    return clean[:30] if clean else "Scope"

def inspect_and_parse(file_bytes, file_name):
    if file_name.lower().endswith('.csv'):
        try:
            df_raw = pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=str)
        except Exception:
            df_raw = pd.read_csv(io.BytesIO(file_bytes), header=None, encoding='cp932', dtype=str)
    elif file_name.lower().endswith('.xls'):
        try:
            df_raw = pd.read_excel(io.BytesIO(file_bytes), header=None, dtype=str, engine='xlrd')
        except Exception:
            df_raw = pd.read_excel(io.BytesIO(file_bytes), header=None, dtype=str)
    else:
        df_raw = pd.read_excel(io.BytesIO(file_bytes), header=None, dtype=str)

    header_row_idx = -1
    col_map = {}

    for r in range(min(25, len(df_raw))):
        row_vals = [clean_str(v).lower() for v in df_raw.iloc[r].tolist()]
        temp_map = {}

        for c, text in enumerate(row_vals):
            if not text:
                continue

            if any(k in text for k in ["expert type", "scope", "part", "category", "タイプ", "スコープ"]):
                if "Scope" not in temp_map: temp_map["Scope"] = c
            elif any(k in text for k in ["name", "expert name", "名前", "氏名", "エキスパート名"]) and "filename" not in text and "thirdbridge" not in text:
                if "Name" not in temp_map: temp_map["Name"] = c
            elif any(k in text for k in ["title", "titles", "役職", "タイトル", "要職"]):
                if "Relevant Titles" not in temp_map: temp_map["Relevant Titles"] = c
            elif any(k in text for k in ["experience", "screening", "経歴要約", "要約", "スクリーニング"]):
                if "Relevant experience" not in temp_map: temp_map["Relevant experience"] = c
            elif any(k in text for k in ["rate", "hourly", "fee", "price", "cost", "単価", "時給", "料金"]):
                if "Hourly Rate" not in temp_map: temp_map["Hourly Rate"] = c
            elif any(k in text for k in ["employment", "history", "career", "経歴", "職歴", "職務経歴"]):
                if "Employment History" not in temp_map: temp_map["Employment History"] = c
            elif any(k in text for k in ["location", "country", "city", "place", "所在地", "国", "場所"]):
                if "Location" not in temp_map: temp_map["Location"] = c
            elif "status" in text or "ステータス" in text:
                if "Status" not in temp_map: temp_map["Status"] = c
            elif any(k in text for k in ["identity", "verification", "id verification", "本人確認"]):
                if "Identity Verification Status" not in temp_map: temp_map["Identity Verification Status"] = c

        if ("Name" in temp_map or "Relevant experience" in temp_map or "Scope" in temp_map) and len(temp_map) >= 2:
            header_row_idx = r
            col_map = temp_map
            break

    if header_row_idx == -1:
        header_row_idx = 0

    number_col_idx = 4 # E列 (index 4)

    has_id = "Identity Verification Status" in col_map

    parsed_rows = []
    has_any_name = False
    
    for r in range(header_row_idx + 1, len(df_raw)):
        row = df_raw.iloc[r]
        
        def get_field(key):
            if key in col_map:
                c = col_map[key]
                if c < len(row):
                    return clean_str(row[c])
            return ""

        name = get_field("Name")
        if name:
            has_any_name = True

        scope = get_field("Scope")
        titles = get_field("Relevant Titles")
        exp = get_field("Relevant experience")
        rate_raw = get_field("Hourly Rate")
        emp_raw = get_field("Employment History")
        loc = get_field("Location")
        status = get_field("Status")
        id_ver = get_field("Identity Verification Status")

        raw_num = clean_str(row[number_col_idx]) if number_col_idx < len(row) else ""
        if "thirdbridge" in raw_num.lower():
            raw_num = ""

        number_val = raw_num
        if raw_num:
            try:
                if "." in raw_num:
                    number_val = float(raw_num)
                else:
                    number_val = int(raw_num)
            except ValueError:
                number_val = raw_num

        cleaned_emp = ""
        if emp_raw:
            emp_lines = [line.strip() for line in re.split(r'[\r\n]+', str(emp_raw)) if line.strip() and line.strip().lower() != "nan"]
            cleaned_emp = "\n".join(emp_lines)

        if not any([name, scope, raw_num, titles, exp, rate_raw, cleaned_emp, loc]):
            continue

        if parsed_rows and not name and not scope and not rate_raw and not raw_num:
            prev = parsed_rows[-1]
            if exp:
                prev["exp"] = (prev["exp"] + "\n" + exp).strip()
            if titles:
                prev["titles"] = (prev["titles"] + "\n" + titles).strip()
            if cleaned_emp:
                prev["emp"] = (prev["emp"] + "\n" + cleaned_emp).strip()
            continue

        if not scope:
            scope = "その他"

        is_consulted = (status.lower() == "consulted")

        rate_val = rate_raw
        if rate_raw:
            num_match = re.search(r'[\d\.]+', rate_raw.replace(',', ''))
            if num_match:
                try:
                    rate_val = float(num_match.group(0))
                except ValueError:
                    rate_val = rate_raw

        id_sym = ""
        if has_id and id_ver:
            if "verified" in id_ver.lower() and "not" not in id_ver.lower():
                id_sym = "◯"
            elif "not verified" in id_ver.lower():
                id_sym = "×"
            else:
                id_sym = id_ver

        parsed_rows.append({
            "scope": scope,
            "number": number_val,
            "name": name,
            "titles": titles,
            "exp": exp,
            "rate": rate_val,
            "emp": cleaned_emp,
            "loc": loc,
            "id_sym": id_sym,
            "is_consulted": is_consulted
        })

    return parsed_rows, has_id, has_any_name

def process_excel(file):
    file_bytes = file.read()
    data_items, has_id, has_any_name = inspect_and_parse(file_bytes, file.name)

    if not data_items:
        raise ValueError("有効なデータ行が見つかりませんでした。ファイルの内容を確認してください。")

    title_val = "ThirdBridge"
    scope_groups = {}

    headers = ["Scope", "Number"]
    if has_any_name:
        headers.append("Name")
    headers.extend(["Relevant Titles", "Relevant experience", "Hourly Rate", "Employment History", "Location"])
    if has_id:
        headers.append("ID Verification")

    formatted_data_list = []
    for item in data_items:
        row_vals = [item["scope"], item["number"]]
        if has_any_name:
            row_vals.append(item["name"])
        row_vals.extend([
            item["titles"],
            item["exp"],
            item["rate"],
            item["emp"],
            item["loc"]
        ])
        if has_id:
            row_vals.append(item["id_sym"])

        entry = {"data": row_vals, "is_consulted": item["is_consulted"]}
        formatted_data_list.append(entry)
        scope_groups.setdefault(item["scope"], []).append(entry)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_regular = Font(name="Meiryo UI", size=9)
    font_bold = Font(name="Meiryo UI", size=9, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    fill_consulted = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_header = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    def build_sheet(ws, sheet_title, items):
        ws.freeze_panes = "D4" if has_any_name else "C4"

        ws["A1"] = sheet_title
        ws["A1"].font = font_bold

        ws["D1"] = "✅ > screening済み"
        ws["D1"].font = font_bold

        ws["D2"] = "黄色セル > コンサル済み"
        ws["D2"].font = font_bold
        ws["D2"].fill = fill_consulted

        for col_num, h_text in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num, value=h_text)
            cell.font = font_bold
            cell.fill = fill_header
            cell.alignment = Alignment(vertical="top")

        for r_idx, item in enumerate(items, 4):
            row_data = item["data"]
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = font_regular

                if item["is_consulted"]:
                    cell.fill = fill_consulted
                else:
                    cell.fill = fill_white

                col_name = headers[c_idx - 1]
                cell.value = val

                if col_name == "Hourly Rate" and isinstance(val, (int, float)):
                    cell.number_format = '$#,##0'
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif col_name in ["Number", "Location", "ID Verification"]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    if col_name == "Number" and isinstance(val, float):
                        cell.number_format = '0.0'

        width_map = {
            "Scope": 22, "Number": 10, "Name": 16,
            "Relevant Titles": 38, "Relevant experience": 60,
            "Hourly Rate": 14, "Employment History": 65,
            "Location": 10, "ID Verification": 14
        }
        for c_i, h_text in enumerate(headers, 1):
            w = width_map.get(h_text, 15)
            ws.column_dimensions[get_column_letter(c_i)].width = w

    # 1. 全員一覧
    ws_all = wb.create_sheet(title="全員一覧")
    build_sheet(ws_all, title_val, formatted_data_list)

    # 2. Scope別シート
    existing_titles = set(["全員一覧"])
    for scope_name, items in scope_groups.items():
        clean_name = sanitize_sheet_name(scope_name)
        
        base_name = clean_name
        counter = 1
        while clean_name in existing_titles:
            clean_name = f"{base_name[:25]}_{counter}"
            counter += 1
            
        existing_titles.add(clean_name)
        ws_scope = wb.create_sheet(title=clean_name)
        build_sheet(ws_scope, f"{title_val} ({scope_name})", items)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

if uploaded_file is not None:
    if st.button("変換してダウンロード", type="primary"):
        with st.spinner("データを正確に処理中..."):
            try:
                processed_data = process_excel(uploaded_file)
                st.success("✨ 変換が完了しました！")
                st.download_button(
                    label="📥 エクセルファイルをダウンロード",
                    data=processed_data,
                    file_name=f"Formatted_{uploaded_file.name}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                st.error(f"エラーが発生しました: {e}")
